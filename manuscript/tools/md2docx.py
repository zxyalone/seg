"""Convert a Chinese manuscript markdown file to .docx with numbered citations.

usage: python md2docx.py <in.md> <refs.bib> <out.docx>

- \\cite{a,b} -> [1, 2] numbered by first appearance; reference list built from refs.bib
  and written under the "## 参考文献" heading.
- $$...$$ and $...$ become native Word equations (pandoc OMML); \\tag{n} -> (n).
- Images are resolved relative to the markdown file.
- Requires pandoc on PATH and the python packages bibtexparser (1.x) and python-docx.
"""
import os
import re
import subprocess
import sys
import tempfile

import bibtexparser
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

src, bibpath, out = sys.argv[1:4]
md = open(src, encoding="utf-8").read()

# ---------- citations: number by first appearance ----------
order = []
for m in re.finditer(r"\\cite\{([^}]*)\}", md):
    for k in m.group(1).split(","):
        k = k.strip()
        if k not in order:
            order.append(k)
num = {k: i + 1 for i, k in enumerate(order)}


def compress(ns):
    ns = sorted(set(ns))
    parts, i = [], 0
    while i < len(ns):
        j = i
        while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1:
            j += 1
        if j - i >= 2:
            parts.append(f"{ns[i]}–{ns[j]}")
        else:
            parts.extend(str(n) for n in ns[i:j + 1])
        i = j + 1
    return "[" + ", ".join(parts) + "]"


md = re.sub(r"\s?\\cite\{([^}]*)\}",
            lambda m: compress([num[k.strip()] for k in m.group(1).split(",")]), md)

# ---------- reference list ----------
with open(bibpath, encoding="utf-8") as f:
    bib = {e["ID"]: e for e in bibtexparser.load(f).entries}


def clean(s):
    s = s.replace('{\\"a}', "ä").replace('{\\"o}', "ö").replace('{\\"u}', "ü")
    s = re.sub(r"\\[a-zA-Z]+\s*", "", s)
    return re.sub(r"[{}]", "", s).replace("--", "–").strip()


def authors(s):
    out = []
    for a in re.split(r"\s+and\s+", clean(s)):
        if a == "others":
            out.append("et al.")
            continue
        if "," in a:
            last, first = [x.strip() for x in a.split(",", 1)]
        else:
            *fs, last = a.split()
            first = " ".join(fs)
        ini = " ".join("-".join(q[0] + "." for q in p.split("-") if q) for p in first.split() if p)
        out.append(f"{ini} {last}".strip())
    if out and out[-1] == "et al.":
        return ", ".join(out[:-1]) + ", et al."
    return ", ".join(out)


refs = []
for k in order:
    e = bib.get(k)
    if e is None:
        refs.append(f"[{num[k]}] [[缺少文献条目：{k}]]")
        continue
    venue = clean(e.get("journal") or e.get("booktitle") or "")
    s = f"[{num[k]}] {authors(e.get('author', ''))}, {clean(e.get('title', ''))}, {venue}"
    if e.get("volume"):
        s += f" {clean(e['volume'])}"
    s += f" ({e.get('year', '')})"
    if e.get("pages"):
        s += f" {clean(e['pages'])}"
    s += "."
    if e.get("doi"):
        s += f" https://doi.org/{e['doi']}"
    elif e.get("note"):
        s += f" {clean(e['note'])}."
    refs.append(s)

md = re.sub(r"## 参考文献\n\n引用键与 `manuscript/refs.bib` 一致。\n?",
            "## 参考文献\n\n" + "\n\n".join(refs) + "\n", md)

# ---------- math: \tag is not supported in OMML ----------
md = re.sub(r"\\tag\{(\d+)\}", r"\\qquad (\1)", md)

with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False, encoding="utf-8") as t:
    t.write(md)
    tmp = t.name
subprocess.run(["pandoc", tmp, "-f", "markdown+tex_math_dollars", "-t", "docx",
                "--resource-path", os.path.dirname(os.path.abspath(src)),
                "--shift-heading-level-by=-1", "-o", out], check=True)
os.unlink(tmp)

# ---------- styling ----------
doc = Document(out)
for sec in doc.sections:
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.2)


def set_font(style, size=None, bold=None, east="宋体", latin="Times New Roman"):
    style.font.name = latin
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi"):
        rfonts.set(qn(a), latin)
    rfonts.set(qn("w:eastAsia"), east)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)


styles = doc.styles
names = [s.name for s in styles]
for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
    if name in names:
        set_font(styles[name], 10.5)
for name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
    if name in names:
        set_font(styles[name], size, True, east="黑体")
if "Image Caption" in names:
    set_font(styles["Image Caption"], 9)
    styles["Image Caption"].font.italic = False

for p in doc.paragraphs:
    if p.style.name in ("Body Text", "First Paragraph"):
        p.paragraph_format.first_line_indent = Pt(21)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(3)
    elif p.style.name in ("Captioned Figure", "Figure"):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
    elif p.style.name == "Image Caption":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(9)
        for r in p.runs:
            r.font.italic = False

# pandoc caps images at its own default text width; stretch wide figures to the A4 text width
text_w = Cm(21 - 2 * 2.0)
for shape in doc.inline_shapes:
    if shape.width > 0.8 * text_w:
        ratio = text_w / shape.width
        shape.width, shape.height = int(text_w), int(shape.height * ratio)


def border_table(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        b.append(el)
    tblPr.append(b)


for tbl in doc.tables:
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    border_table(tbl)
    ncol = len(tbl.columns)
    size = 10 if ncol <= 4 else 8 if ncol <= 9 else 7
    for r_i, row in enumerate(tbl.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = None
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(size)
                    if r_i == 0:
                        run.font.bold = True

doc.save(out)
print(f"{len(order)} citations, {len(doc.tables)} tables, {len(doc.inline_shapes)} images -> {out}")
